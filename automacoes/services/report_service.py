from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def gerar_relatorio_xlsx(registros: list[dict[str, Any]], destino: str | Path) -> str:
    """Exporta os registros processados para uma planilha XLSX formatada."""
    if not registros:
        raise ValueError("Não há registros para exportar.")

    output = Path(destino).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    colunas: list[str] = []
    for registro in registros:
        if isinstance(registro, dict):
            for chave in registro:
                if chave not in colunas:
                    colunas.append(str(chave))
    if not colunas:
        raise ValueError("Os dados recebidos não têm campos exportáveis.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório API"
    ws.append(colunas)
    for registro in registros:
        row = []
        for chave in colunas:
            valor = registro.get(chave, "") if isinstance(registro, dict) else ""
            if isinstance(valor, (dict, list, tuple)):
                valor = json.dumps(valor, ensure_ascii=False)
            row.append("" if valor is None else str(valor))
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, col in enumerate(ws.columns, start=1):
        longest = max((len(str(cell.value or "")) for cell in list(col)[:200]), default=10)
        ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 42)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(output)
    return str(output)
